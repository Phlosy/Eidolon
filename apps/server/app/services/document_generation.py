"""Template-driven Markdown, DOCX and PPTX generation over the shared Drive."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from pptx.util import Pt as PptPt
from sqlalchemy.orm import Session

from app.models.enums import DocumentCategory
from app.models.project import Project
from app.models.project_delivery import DocumentArtifact, ProjectPhase
from app.repositories import drive as drive_repo
from app.repositories import project_delivery as delivery_repo
from app.services import drive as drive_service


@dataclass(frozen=True)
class DocumentSection:
    title: str
    body: str
    bullets: list[str]


@dataclass(frozen=True)
class DocumentSource:
    title: str
    summary: str
    sections: list[DocumentSection]


class DocumentGenerationService:
    """One generation seam for formal project documents and review material."""

    def generate_markdown(self, source: DocumentSource) -> str:
        chunks = [f"# {source.title}", "", source.summary]
        for section in source.sections:
            chunks.extend(["", f"## {section.title}", "", section.body])
            chunks.extend(f"- {item}" for item in section.bullets)
        return "\n".join(chunks).strip() + "\n"

    def generate_docx(self, source: DocumentSource) -> bytes:
        document = Document()
        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10.5)
        title = document.add_heading(source.title, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        summary = document.add_paragraph(source.summary)
        summary.style = styles["Subtitle"]
        for section in source.sections:
            document.add_heading(section.title, level=1)
            if section.body:
                document.add_paragraph(section.body)
            for item in section.bullets:
                document.add_paragraph(item, style="List Bullet")
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def generate_pptx(self, source: DocumentSource) -> bytes:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        title_slide.shapes.title.text = source.title
        title_slide.placeholders[1].text = source.summary
        self._style_slide(title_slide)
        for section in source.sections[:8]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = section.title
            frame = slide.placeholders[1].text_frame
            frame.clear()
            points = section.bullets or ([section.body] if section.body else [])
            for index, point in enumerate(points[:5]):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = point
                paragraph.font.size = PptPt(22)
                paragraph.space_after = PptPt(12)
            self._style_slide(slide)
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()

    def generate_speaker_notes(self, source: DocumentSource) -> str:
        lines = [f"# {source.title} · Speaker Notes", ""]
        lines.extend(["## Slide 01 · Opening", "", source.summary, ""])
        for index, section in enumerate(source.sections[:8], start=2):
            lines.extend(
                [
                    f"## Slide {index:02d} · {section.title}",
                    "",
                    section.body or f"本页说明{section.title}的关键结论与决策依据。",
                    "",
                    "讲解要点：",
                    *(f"- {item}" for item in section.bullets[:5]),
                    "",
                ]
            )
        return "\n".join(lines).strip() + "\n"

    def persist(
        self,
        db: Session,
        *,
        project: Project,
        phase: ProjectPhase,
        category: str,
        document_type: str,
        title: str,
        format: str,
        content: str | bytes,
        author_employee_id: int | None,
        source_document_id: int | None = None,
        review_id: int | None = None,
        metadata: dict | None = None,
    ) -> DocumentArtifact:
        folders = drive_service.ensure_project_folders(db, project)
        directory = {
            DocumentCategory.source.value: "source",
            DocumentCategory.build.value: "release",
            DocumentCategory.delivery.value: "release",
            DocumentCategory.test.value: "tests",
        }.get(category, "docs")
        parent = folders[directory]
        minor = delivery_repo.next_document_minor(db, project.id, document_type)
        version_label = f"v0.{minor}"
        # The first generated copy keeps the clean business-facing title. Later
        # review attempts surface their artifact revision in Drive so two files
        # with the same title never look like accidental duplicates.
        display_name = title if minor == 1 else f"{title} {version_label}"
        filename = f"{display_name}.{format}"
        if format == "markdown":
            node = drive_service.create_document(
                db,
                parent=parent,
                name=display_name,
                content=str(content),
                doc_type=document_type,
                project_id=project.id,
                owner_employee_id=author_employee_id,
                message=f"Generate {document_type} {version_label}",
            )
        else:
            if not isinstance(content, bytes):
                raise TypeError("binary document content must be bytes")
            node = drive_service.create_binary_document(
                db,
                parent=parent,
                name=filename,
                content=content,
                extension=format,
                doc_type=document_type,
                project_id=project.id,
                owner_employee_id=author_employee_id,
                message=f"Generate {document_type} {version_label}",
            )
        revision = drive_repo.list_revisions(db, node.id)[0]
        return delivery_repo.create_document(
            db,
            project_id=project.id,
            phase_id=phase.id,
            review_id=review_id,
            category=category,
            document_type=document_type,
            title=title,
            format=format,
            version_major=0,
            version_minor=minor,
            version_label=version_label,
            drive_node_id=node.id,
            drive_revision_id=revision.id,
            author_employee_id=author_employee_id,
            review_status="review" if review_id else "draft",
            baseline_status="none",
            source_document_id=source_document_id,
            metadata_json=metadata or {},
        )

    @staticmethod
    def content_hash(document_ids: list[int]) -> str:
        return hashlib.sha256(",".join(map(str, document_ids)).encode()).hexdigest()

    @staticmethod
    def _style_slide(slide) -> None:
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(15, 23, 42)
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.color.rgb = RGBColor(241, 245, 249)


document_generation_service = DocumentGenerationService()
